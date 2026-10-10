"""PDF / DOCX export built fully in memory (serverless hosts have no writable disk)."""
import io
import re
from xml.sax.saxutils import escape

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _plain(md):
    md = re.sub(r"^#{1,6}\s*", "", md, flags=re.M)
    md = _BOLD.sub(r"\1", md)
    return re.sub(r"`([^`]*)`", r"\1", md)


def note_pdf(title, markdown):
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    styles = getSampleStyleSheet()
    styles["Title"].alignment = TA_CENTER
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 6 * mm)]
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 3 * mm))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        text = _BOLD.sub(r"<b>\1</b>", escape(m.group(2) if m else stripped))
        if m:
            story.append(Paragraph(text, styles["Heading2" if len(m.group(1)) < 3 else "Heading3"]))
        else:
            story.append(Paragraph(re.sub(r"^[-*]\s+", "&bull; ", text), styles["BodyText"]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title=title).build(story)
    return buf.getvalue()


def note_docx(title, markdown):
    from docx import Document
    doc = Document()
    doc.add_heading(title, 0)
    for line in _plain(markdown).splitlines():
        if line.strip():
            doc.add_paragraph(line.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def plan_pdf(title, plan):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 6 * mm)]
    for day in plan:
        story += [Paragraph(f"Day {day['day']} - {day['date']}: {escape(day['title'])}", styles["Heading3"]),
                  Paragraph(f"Study time: {day['minutes']} minutes", styles["BodyText"])]
        for t in day["tasks"]:
            story.append(Paragraph(("[x] " if t.get("done") else "[ ] ") + escape(t["text"]), styles["BodyText"]))
        story.append(Spacer(1, 4 * mm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title=title).build(story)
    return buf.getvalue()
